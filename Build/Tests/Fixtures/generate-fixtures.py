#!/usr/bin/env python3
"""
Generates the round-trip fixtures in this folder.

    python3 Build/Tests/Fixtures/generate-fixtures.py

Needs python-docx (pip install python-docx). The fixtures are committed; run
this only to change them. python-docx starts from Word's own default template,
so the packages carry real Word parts (styles, theme, settings, numbering,
custom XML) next to the content added here:

- features.docx: headings 1-4, run formatting, a hyperlink, bulleted and
  numbered lists (two levels), a table with merged cells and a repeated header
  row, an image, headers and footers (first page and default, PAGE field), a
  second landscape section and a comment.
- structure.docx: content controls (block, inline, nested, locked, drop-down,
  date, checkbox, data-bound), bookmarks with an internal link, tracked
  insertions and deletions, a footnote, an extra custom XML part and custom
  document properties.
- localized-styles.docx: style IDs the way a German Word writes them
  (Standard, berschrift1, berschrift2) with Heading 3 and 4 left latent.
"""

import io
import re
import struct
import zipfile
import zlib
from pathlib import Path

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def png(width: int, height: int) -> bytes:
    """A small RGB gradient PNG, written by hand so no imaging library is needed."""
    rows = b''
    for y in range(height):
        rows += b'\x00' + b''.join(
            bytes((x * 255 // max(width - 1, 1), y * 255 // max(height - 1, 1), 160)) for x in range(width)
        )

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data) & 0xFFFFFFFF)

    header = struct.pack('>IIBBBBB', width, height, 8, 2, 0, 0, 0)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', header) + chunk(b'IDAT', zlib.compress(rows, 9)) + chunk(b'IEND', b'')


def add_field(paragraph, instruction: str, cached: str) -> None:
    """A complex field (begin / instrText / separate / result / end)."""
    for xml in (
        f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>',
        f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> {instruction} </w:instrText></w:r>',
        f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>',
        f'<w:r {nsdecls("w")}><w:t>{cached}</w:t></w:r>',
        f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>',
    ):
        paragraph._p.append(parse_xml(xml))


def add_hyperlink(paragraph, url: str, text: str) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    paragraph._p.append(
        parse_xml(
            f'<w:hyperlink {nsdecls("w", "r")} r:id="{r_id}" w:history="1">'
            '<w:r><w:rPr><w:color w:val="0563C1"/><w:u w:val="single"/></w:rPr>'
            f'<w:t>{text}</w:t></w:r></w:hyperlink>'
        )
    )


def features() -> None:
    document = docx.Document()
    document.core_properties.title = 'Round-trip fixture: features'
    document.core_properties.author = 'docx_editor test suite'

    document.add_paragraph('Round-trip features', style='Title')
    for level in range(1, 5):
        document.add_heading(f'Heading level {level}', level=level)

    body = document.add_paragraph('Plain text, ')
    body.add_run('bold').bold = True
    body.add_run(', ')
    body.add_run('italic').italic = True
    body.add_run(', ')
    body.add_run('underlined').underline = True
    body.add_run(', ')
    coloured = body.add_run('coloured 14pt')
    coloured.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    coloured.font.size = Pt(14)
    body.add_run(' and a link to ')
    add_hyperlink(body, 'https://typo3.org/', 'typo3.org')
    body.add_run('.')

    commented = document.add_paragraph('This sentence carries a comment.')
    document.add_comment(commented.runs, text='Please check this sentence.', author='Reviewer', initials='RV')

    document.add_paragraph('First bullet', style='List Bullet')
    document.add_paragraph('Nested bullet', style='List Bullet 2')
    document.add_paragraph('Second bullet', style='List Bullet')
    document.add_paragraph('First step', style='List Number')
    document.add_paragraph('Second step', style='List Number')
    document.add_paragraph('Third step', style='List Number')

    table = document.add_table(rows=4, cols=3)
    table.style = 'Table Grid'
    for column, title in enumerate(('Name', 'Role', 'Notes')):
        table.cell(0, column).text = title
    table.rows[0]._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    table.cell(1, 0).text = 'Ada'
    table.cell(1, 1).text = 'Engineer'
    table.cell(2, 0).merge(table.cell(2, 1)).text = 'Merged across two columns'
    table.cell(1, 2).merge(table.cell(3, 2)).text = 'Merged down three rows'
    table.cell(3, 0).text = 'Grace'
    table.cell(3, 1).text = 'Admiral'

    document.add_paragraph('An image follows.')
    document.add_picture(io.BytesIO(png(64, 32)), width=Inches(1.5))

    section = document.sections[0]
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = 'First page header'
    section.header.paragraphs[0].text = 'Running header'
    footer = section.footer.paragraphs[0]
    footer.text = 'Page '
    add_field(footer, 'PAGE', '2')

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph('Second page, still portrait.')

    landscape = document.add_section()
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width, landscape.page_height = landscape.page_height, landscape.page_width
    document.add_paragraph('A landscape section.')

    document.save(HERE / 'features.docx')


CUSTOM_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<webcon:page xmlns:webcon="urn:webconsulting:typo3:page" uid="42" language="0">'
    '<webcon:title>Round-trip page</webcon:title><webcon:content uid="7" CType="text"/>'
    '</webcon:page>'
)
CUSTOM_XML_PROPS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="no"?>'
    '<ds:datastoreItem ds:itemID="{8B7C4A1E-2F3D-4C5B-9A6E-1D2C3B4A5F60}" '
    'xmlns:ds="http://schemas.openxmlformats.org/officeDocument/2006/customXml">'
    '<ds:schemaRefs><ds:schemaRef ds:uri="urn:webconsulting:typo3:page"/></ds:schemaRefs>'
    '</ds:datastoreItem>'
)
CUSTOM_PROPERTIES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="typo3PageUid">'
    '<vt:i4>42</vt:i4></property>'
    '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="3" name="typo3Site">'
    '<vt:lpwstr>main</vt:lpwstr></property>'
    '</Properties>'
)
FOOTNOTES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    f'<w:footnotes xmlns:w="{W_NS}">'
    '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
    '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
    '<w:footnote w:id="1"><w:p><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r>'
    '<w:r><w:t xml:space="preserve"> A footnote that has to survive.</w:t></w:r></w:p></w:footnote>'
    '</w:footnotes>'
)


def sdt_block(tag: str, alias: str, sdt_id: int, inner: str, extra_pr: str = '') -> str:
    return (
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:alias w:val="{alias}"/><w:tag w:val="{tag}"/>'
        f'<w:id w:val="{sdt_id}"/>{extra_pr}</w:sdtPr><w:sdtContent>{inner}</w:sdtContent></w:sdt>'
    )


def structure() -> None:
    document = docx.Document()
    document.core_properties.title = 'Round-trip fixture: structure'
    document.add_heading('Structure that must survive', level=1)
    body = document.element.body
    sect_pr = body.find(qn('w:sectPr'))

    def append(xml: str) -> None:
        sect_pr.addprevious(parse_xml(xml))

    append(
        sdt_block(
            'typo3:content:7',
            'Content element 7',
            1001,
            '<w:p><w:r><w:t>Block content control, locked against deletion.</w:t></w:r></w:p>',
            '<w:lock w:val="sdtLocked"/>',
        )
    )
    append(
        sdt_block(
            'typo3:container',
            'Container',
            1002,
            '<w:p><w:r><w:t xml:space="preserve">Outer control with </w:t></w:r>'
            '<w:sdt><w:sdtPr><w:alias w:val="Inline"/><w:tag w:val="typo3:field:header"/><w:id w:val="1003"/>'
            '<w:lock w:val="contentLocked"/></w:sdtPr>'
            '<w:sdtContent><w:r><w:t>a nested inline control</w:t></w:r></w:sdtContent></w:sdt>'
            '<w:r><w:t>.</w:t></w:r></w:p>',
        )
    )
    append(
        '<w:p ' + nsdecls('w') + '><w:r><w:t xml:space="preserve">Choose: </w:t></w:r>'
        '<w:sdt><w:sdtPr><w:alias w:val="Choice"/><w:tag w:val="typo3:choice"/><w:id w:val="1004"/>'
        '<w:dropDownList><w:listItem w:displayText="First" w:value="1"/>'
        '<w:listItem w:displayText="Second" w:value="2"/></w:dropDownList></w:sdtPr>'
        '<w:sdtContent><w:r><w:t>First</w:t></w:r></w:sdtContent></w:sdt>'
        '<w:r><w:t xml:space="preserve">, date: </w:t></w:r>'
        '<w:sdt><w:sdtPr><w:alias w:val="Date"/><w:tag w:val="typo3:date"/><w:id w:val="1005"/>'
        '<w:date w:fullDate="2026-09-23T00:00:00Z"><w:dateFormat w:val="dd.MM.yyyy"/><w:lid w:val="de-AT"/>'
        '<w:storeMappedDataAs w:val="dateTime"/><w:calendar w:val="gregorian"/></w:date></w:sdtPr>'
        '<w:sdtContent><w:r><w:t>23.09.2026</w:t></w:r></w:sdtContent></w:sdt>'
        '<w:r><w:t xml:space="preserve">, done: </w:t></w:r>'
        '<w:sdt><w:sdtPr><w:alias w:val="Done"/><w:tag w:val="typo3:done"/><w:id w:val="1006"/>'
        '<w14:checkbox xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml">'
        '<w14:checked w14:val="1"/><w14:checkedState w14:val="2612" w14:font="MS Gothic"/>'
        '<w14:uncheckedState w14:val="2610" w14:font="MS Gothic"/></w14:checkbox></w:sdtPr>'
        '<w:sdtContent><w:r><w:t>☒</w:t></w:r></w:sdtContent></w:sdt></w:p>'
    )
    append(
        sdt_block(
            'typo3:bound:title',
            'Bound title',
            1007,
            '<w:p><w:r><w:t>Round-trip page</w:t></w:r></w:p>',
            '<w:dataBinding w:prefixMappings="xmlns:webcon=\'urn:webconsulting:typo3:page\'" '
            'w:xpath="/webcon:page[1]/webcon:title[1]" w:storeItemID="{8B7C4A1E-2F3D-4C5B-9A6E-1D2C3B4A5F60}"/>',
        )
    )
    append(
        '<w:p ' + nsdecls('w') + '><w:bookmarkStart w:id="10" w:name="intro_range"/>'
        '<w:r><w:t xml:space="preserve">A bookmark spans </w:t></w:r><w:r><w:rPr><w:b/></w:rPr><w:t>two runs</w:t></w:r>'
        '<w:bookmarkEnd w:id="10"/><w:r><w:t xml:space="preserve"> and </w:t></w:r>'
        '<w:hyperlink w:anchor="intro_range"><w:r><w:rPr><w:u w:val="single"/></w:rPr><w:t>links back</w:t></w:r></w:hyperlink>'
        '<w:r><w:t>.</w:t></w:r><w:bookmarkStart w:id="11" w:name="_GoBack"/><w:bookmarkEnd w:id="11"/></w:p>'
    )
    append(
        '<w:p ' + nsdecls('w') + '><w:r><w:t xml:space="preserve">Tracked: </w:t></w:r>'
        '<w:ins w:id="20" w:author="Editor A" w:date="2026-09-20T10:00:00Z"><w:r><w:t>inserted words</w:t></w:r></w:ins>'
        '<w:r><w:t xml:space="preserve"> and </w:t></w:r>'
        '<w:del w:id="21" w:author="Editor B" w:date="2026-09-21T11:00:00Z"><w:r><w:delText>deleted words</w:delText></w:r></w:del>'
        '<w:r><w:t>.</w:t></w:r></w:p>'
    )
    append(
        '<w:p ' + nsdecls('w') + '><w:r><w:t>A sentence with a footnote</w:t></w:r>'
        '<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteReference w:id="1"/></w:r>'
        '<w:r><w:t>.</w:t></w:r></w:p>'
    )

    raw = io.BytesIO()
    document.save(raw)
    add_package_parts(raw.getvalue(), HERE / 'structure.docx')


def add_package_parts(data: bytes, target: Path) -> None:
    """Footnotes, a second custom XML item and custom properties, wired into the package."""
    source = zipfile.ZipFile(io.BytesIO(data))
    parts = {name: source.read(name) for name in source.namelist()}

    types = parts['[Content_Types].xml'].decode()
    types = types.replace(
        '</Types>',
        '<Override PartName="/word/footnotes.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
        '<Override PartName="/customXml/itemProps2.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.customXmlProperties+xml"/>'
        '<Override PartName="/docProps/custom.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/>'
        '</Types>',
    )
    parts['[Content_Types].xml'] = types.encode()

    root_rels = parts['_rels/.rels'].decode()
    root_rels = root_rels.replace(
        '</Relationships>',
        '<Relationship Id="rIdCustomProps" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" '
        'Target="docProps/custom.xml"/></Relationships>',
    )
    parts['_rels/.rels'] = root_rels.encode()

    document_rels = parts['word/_rels/document.xml.rels'].decode()
    document_rels = document_rels.replace(
        '</Relationships>',
        '<Relationship Id="rIdFootnotes" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>'
        '<Relationship Id="rIdCustomXml2" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml" '
        'Target="../customXml/item2.xml"/></Relationships>',
    )
    parts['word/_rels/document.xml.rels'] = document_rels.encode()

    parts['word/footnotes.xml'] = FOOTNOTES.encode()
    parts['customXml/item2.xml'] = CUSTOM_XML.encode()
    parts['customXml/itemProps2.xml'] = CUSTOM_XML_PROPS.encode()
    parts['customXml/_rels/item2.xml.rels'] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps" '
        'Target="itemProps2.xml"/></Relationships>'
    ).encode()
    parts['docProps/custom.xml'] = CUSTOM_PROPERTIES.encode()

    write_package(parts, target)


def write_package(parts: dict, target: Path) -> None:
    with zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as out:
        out.writestr('[Content_Types].xml', parts.pop('[Content_Types].xml'))
        for name, content in parts.items():
            out.writestr(name, content)


def localized_styles() -> None:
    document = docx.Document()
    document.add_heading('Überschrift der ersten Ebene', level=1)
    document.add_paragraph('Ein Absatz im Standard-Format.')
    document.add_heading('Überschrift der zweiten Ebene', level=2)
    document.add_paragraph('Noch ein Absatz.')

    styles = document.styles.element
    renames = {'Normal': 'Standard', 'Heading1': 'berschrift1', 'Heading2': 'berschrift2'}
    latent = {'Heading3', 'Heading4', 'Heading3Char', 'Heading4Char'}
    for style in list(styles.findall(qn('w:style'))):
        style_id = style.get(qn('w:styleId'))
        if style_id in latent:
            styles.remove(style)
            continue
        if style_id in renames:
            style.set(qn('w:styleId'), renames[style_id])
        for reference in ('w:basedOn', 'w:next', 'w:link'):
            node = style.find(qn(reference))
            if node is None:
                continue
            value = node.get(qn('w:val'))
            if value in latent:
                style.remove(node)
            elif value in renames:
                node.set(qn('w:val'), renames[value])
    for paragraph_style in document.element.body.iter(qn('w:pStyle')):
        value = paragraph_style.get(qn('w:val'))
        if value in renames:
            paragraph_style.set(qn('w:val'), renames[value])

    raw = io.BytesIO()
    document.save(raw)
    source = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    parts = {name: source.read(name) for name in source.namelist()}
    # stylesWithEffects.xml is Word 2010's copy of styles.xml; keep the package consistent.
    parts.pop('word/stylesWithEffects.xml', None)
    parts['[Content_Types].xml'] = (
        parts['[Content_Types].xml']
        .decode()
        .replace(
            '<Override PartName="/word/stylesWithEffects.xml" '
            'ContentType="application/vnd.ms-word.stylesWithEffects+xml"/>',
            '',
        )
        .encode()
    )
    rels = parts['word/_rels/document.xml.rels'].decode()
    rels = re.sub(r'<Relationship [^>]*stylesWithEffects[^>]*/>', '', rels)
    parts['word/_rels/document.xml.rels'] = rels.encode()
    write_package(parts, HERE / 'localized-styles.docx')


if __name__ == '__main__':
    features()
    structure()
    localized_styles()
    for name in ('features.docx', 'structure.docx', 'localized-styles.docx'):
        print(name, (HERE / name).stat().st_size, 'bytes')
